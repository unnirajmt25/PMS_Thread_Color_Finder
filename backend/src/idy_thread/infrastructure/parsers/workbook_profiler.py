"""The workbook profiler (section 5 of the spec).

Structural inspection ONLY — this never writes to a database and never
commits to a field mapping. `profile_rows` is a pure function (plain
Python lists in, a SheetProfile out) so it's unit-testable without any
real .xlsx fixtures; `profile_workbook` is the thin I/O wrapper that reads
a file with openpyxl and calls it per sheet.
"""

from __future__ import annotations

import re
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from ...security.file_validation import sha256_of, validate_file
from .field_heuristics import suggest_field_mapping
from .profile_models import ColumnProfile, ProfileReport, QualityWarning, SheetProfile, WorkbookProfile

MAX_HEADER_SCAN_ROWS = 5
MAX_SAMPLE_VALUES = 5


def derive_likely_source(file_stem: str) -> str:
    """Best-effort source name from the filename alone. This is a
    heuristic, not a database lookup — see docs/DATABASE.md for why
    Source/Manufacturer/Brand must NOT be conflated even though this is
    the only signal the profiler has for "source"."""
    return re.sub(r"\s*Thread\s*Charts?\s*$", "", file_stem, flags=re.IGNORECASE).strip() or file_stem


def _cell_str(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _infer_column_type(values: list[Any]) -> str:
    non_empty = [v for v in values if _cell_str(v) != ""]
    if not non_empty:
        return "empty"

    numeric_count = 0
    for v in non_empty:
        if isinstance(v, (int, float)):
            numeric_count += 1
            continue
        s = _cell_str(v)
        try:
            float(s)
            numeric_count += 1
        except ValueError:
            pass

    if numeric_count == len(non_empty):
        return "integer" if all(float(_cell_str(v) if not isinstance(v, (int, float)) else v).is_integer()
                                 for v in non_empty) else "decimal"
    if numeric_count == 0:
        return "text"
    return "mixed"


def _detect_header_row(rows: list[list[Any]]) -> int:
    """Scans the first few rows for the one that looks most like a header
    (mostly non-empty text cells, not numeric data). Falls back to row 1
    (index 0) if nothing else stands out — true for every current sample
    file, but NOT hard-coded: a differently-shaped future file can still
    be detected correctly here."""
    best_idx = 0
    best_score = -1.0
    for idx, row in enumerate(rows[:MAX_HEADER_SCAN_ROWS]):
        cells = [_cell_str(c) for c in row]
        non_empty = [c for c in cells if c]
        if len(non_empty) < 2:
            continue
        text_like = sum(1 for c in non_empty if not _looks_numeric(c))
        score = text_like / len(non_empty)
        if score > best_score:
            best_score = score
            best_idx = idx
    return best_idx


def _looks_numeric(s: str) -> bool:
    try:
        float(s)
        return True
    except ValueError:
        return False


def profile_rows(sheet_name: str, rows: list[list[Any]]) -> SheetProfile | None:
    """Pure profiling logic over an already-loaded array-of-arrays sheet.
    Returns None for a genuinely empty sheet."""
    if not rows:
        return None

    header_row_idx = _detect_header_row(rows)
    header_row = rows[header_row_idx]
    headers = [_cell_str(h) for h in header_row]
    column_count = len(headers)

    data_rows = rows[header_row_idx + 1:]
    row_count = len(data_rows)

    # Duplicate headers (case-insensitive, non-blank only)
    seen_headers: dict = {}
    duplicate_headers: list[str] = []
    for h in headers:
        if not h:
            continue
        key = h.lower()
        seen_headers[key] = seen_headers.get(key, 0) + 1
    duplicate_headers = [h for h in headers if h and seen_headers.get(h.lower(), 0) > 1]
    duplicate_headers = sorted(set(duplicate_headers))

    warnings: list[QualityWarning] = []
    if duplicate_headers:
        warnings.append(QualityWarning(
            severity="WARNING", issue_type="duplicate_headers", sheet_name=sheet_name,
            message=f"Duplicate header(s): {', '.join(duplicate_headers)}",
        ))

    # Column-by-column profiling
    columns: list[ColumnProfile] = []
    empty_column_count = 0
    for col_idx in range(column_count):
        col_values = [row[col_idx] if col_idx < len(row) else None for row in data_rows]
        non_empty_values = [v for v in col_values if _cell_str(v) != ""]
        inferred_type = _infer_column_type(col_values)
        if inferred_type == "empty":
            empty_column_count += 1
        header = headers[col_idx]
        if not header and non_empty_values:
            warnings.append(QualityWarning(
                severity="WARNING", issue_type="unlabeled_column_with_data", sheet_name=sheet_name,
                column_name=f"column_{col_idx + 1}",
                message=f"Column {col_idx + 1} has {len(non_empty_values)} data value(s) but no header",
            ))
        columns.append(ColumnProfile(
            index=col_idx,
            header=header,
            inferred_type=inferred_type,
            non_empty_count=len(non_empty_values),
            empty_count=len(col_values) - len(non_empty_values),
            sample_values=[_cell_str(v) for v in non_empty_values[:MAX_SAMPLE_VALUES]],
        ))

    # Empty rows
    empty_row_count = sum(1 for row in data_rows if not any(_cell_str(c) for c in row))

    # Duplicate records: exact full-row duplicates (conservative — this is
    # NOT "same thread code twice", which is a domain-level question for
    # later; here it's purely "identical row appears more than once").
    row_signatures: dict = {}
    for row in data_rows:
        sig = tuple(_cell_str(c) for c in row)
        if all(c == "" for c in sig):
            continue
        row_signatures[sig] = row_signatures.get(sig, 0) + 1
    duplicate_record_count = sum(count - 1 for count in row_signatures.values() if count > 1)
    if duplicate_record_count:
        warnings.append(QualityWarning(
            severity="WARNING", issue_type="duplicate_records", sheet_name=sheet_name,
            message=f"{duplicate_record_count} duplicate row(s) found",
        ))

    # Field mapping suggestions + RGB range validation where we're
    # confident enough to have found R/G/B columns.
    field_mappings = []
    rgb_columns: dict = {}
    for col in columns:
        if not col.header:
            continue
        suggestion = suggest_field_mapping(col.header)
        if suggestion:
            field_mappings.append(suggestion)
            if suggestion.canonical_field in ("red", "green", "blue") and suggestion.confidence in (
                "HIGH", "MEDIUM"
            ):
                rgb_columns[suggestion.canonical_field] = col.index

    if len(rgb_columns) == 3:
        invalid_rgb = 0
        for row in data_rows:
            for idx in rgb_columns.values():
                if idx >= len(row):
                    continue
                s = _cell_str(row[idx])
                if not s:
                    continue
                try:
                    n = float(s)
                    if not (0 <= n <= 255):
                        invalid_rgb += 1
                except ValueError:
                    invalid_rgb += 1
        if invalid_rgb:
            warnings.append(QualityWarning(
                severity="ERROR", issue_type="invalid_rgb_values", sheet_name=sheet_name,
                message=f"{invalid_rgb} out-of-range or non-numeric R/G/B value(s)",
            ))

    return SheetProfile(
        name=sheet_name,
        row_count=row_count,
        column_count=column_count,
        header_row_number=header_row_idx + 1,
        headers=headers,
        duplicate_headers=duplicate_headers,
        empty_row_count=empty_row_count,
        empty_column_count=empty_column_count,
        duplicate_record_count=duplicate_record_count,
        columns=columns,
        field_mappings=field_mappings,
        warnings=warnings,
    )


def profile_workbook(path: Path) -> WorkbookProfile:
    validate_file(path)
    checksum = sha256_of(path)
    size = path.stat().st_size
    likely_source = derive_likely_source(path.stem)

    workbook_warnings: list[QualityWarning] = []
    sheets: list[SheetProfile] = []

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            rows = [list(row) for row in ws.iter_rows(values_only=True)]
            sheet_profile = profile_rows(sheet_name, rows)
            if sheet_profile is None:
                workbook_warnings.append(QualityWarning(
                    severity="INFO", issue_type="empty_sheet", sheet_name=sheet_name,
                    message="Sheet has no data",
                ))
                continue
            sheets.append(sheet_profile)
    finally:
        wb.close()

    if not sheets:
        workbook_warnings.append(QualityWarning(
            severity="ERROR", issue_type="no_usable_sheets",
            message="No sheet in this workbook had any data",
        ))

    return WorkbookProfile(
        file_name=path.name,
        file_path=str(path),
        file_size_bytes=size,
        sha256=checksum,
        likely_source=likely_source,
        sheets=sheets,
        warnings=workbook_warnings,
    )


def build_profile_report(workbooks: list[WorkbookProfile]) -> ProfileReport:
    """Aggregates per-workbook profiles into the overall report — the
    JSON/Markdown/Rich renderers all consume this one object."""
    total_sheets = 0
    total_rows = 0
    warnings_by_severity: dict = {}
    discovered_headers: set = set()

    def _tally(w: QualityWarning) -> None:
        warnings_by_severity[w.severity] = warnings_by_severity.get(w.severity, 0) + 1

    for wb in workbooks:
        for w in wb.warnings:
            _tally(w)
        for sheet in wb.sheets:
            total_sheets += 1
            total_rows += sheet.row_count
            for w in sheet.warnings:
                _tally(w)
            discovered_headers.update(h for h in sheet.headers if h)

    discovered_sources = sorted({wb.likely_source for wb in workbooks})
    total_warnings = sum(warnings_by_severity.values())

    unknown = [
        ("Manufacturer/brand/product-line identity is not reliably derivable from filename + "
        "sheet name alone — every 'likely_source' here is a FILENAME-based guess, not a "
        "verified Manufacturer/Brand/Source relationship. REQUIRES REVIEW before Milestone 3."),
        ("Whether 'Thread Chart' sheet-name column (e.g. 'RA SSR 9') denotes a Brand, a "
        "ProductLine, or neither is REQUIRES REVIEW — see docs/DATABASE.md."),
    ]

    return ProfileReport(
        generated_at=datetime.now(timezone.utc),
        files_analyzed=len(workbooks),
        total_sheets=total_sheets,
        total_rows=total_rows,
        total_warnings=total_warnings,
        warnings_by_severity=warnings_by_severity,
        discovered_sources=discovered_sources,
        discovered_headers=sorted(discovered_headers),
        workbooks=workbooks,
        unknown=unknown,
    )
