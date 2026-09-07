"""Parses the vendor "Thread Chart" workbooks into a flat list of records.

Mirrors the logic in scripts/generate-thread-data.mjs (the JS generator
used by the Thred Finder web app) so the two stay consistent, but reads
the .xlsx files directly with openpyxl instead of going through the
generated JSON.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, asdict
from datetime import date
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook

_NUMERIC_PMS_RE = re.compile(r"^(\d+)\s*C?$", re.IGNORECASE)
_TRAILING_C_RE = re.compile(r"([a-zA-Z0-9])C$")


@dataclass(frozen=True)
class Record:
    vendor: str
    thread_brand: str
    thread_code: str
    thread_color_name: str
    thread_hex: str
    pms_code: str
    pms_name: str
    pms_hex: str
    updated_at: str

    def as_dict(self) -> dict:
        return asdict(self)


def _vendor_name_from_file(path: Path) -> str:
    return path.stem.strip()


def _to_hex(r, g, b) -> str:
    try:
        r, g, b = int(round(float(r))), int(round(float(g))), int(round(float(b)))
    except (TypeError, ValueError):
        return ""
    r, g, b = (max(0, min(255, v)) for v in (r, g, b))
    return f"#{r:02x}{g:02x}{b:02x}"


def _normalize_pms(raw) -> tuple[str, str]:
    if raw is None:
        return "", ""
    trimmed = str(raw).strip()
    if not trimmed or trimmed.lower() == "x":
        return "", ""

    numeric_match = _NUMERIC_PMS_RE.match(trimmed)
    if numeric_match:
        code = numeric_match.group(1)
        return code, f"PMS {code} C"

    spaced = _TRAILING_C_RE.sub(r"\1 C", trimmed)
    spaced = re.sub(r"\s+", " ", spaced).strip()
    return spaced, f"PMS {spaced}"


def _header_index_map(header_row: Iterable) -> dict[str, int]:
    index_map: dict[str, int] = {}
    for i, cell in enumerate(header_row):
        key = str(cell or "").strip().lower()
        if key:
            index_map[key] = i
    return index_map


def _col(index_map: dict[str, int], *names: str) -> int:
    for name in names:
        idx = index_map.get(name.lower())
        if idx is not None:
            return idx
    return -1


def load_records(data_dir: str | Path, *, updated_at: str | None = None) -> list[Record]:
    """Parses every .xlsx workbook in `data_dir` into a flat, deduplicated
    list of Records. Skips Excel's own temp lock files (~$...)."""
    data_dir = Path(data_dir)
    updated_at = updated_at or date.today().isoformat()

    files = sorted(
        p
        for p in data_dir.glob("*.xlsx")
        if not p.name.startswith("~$")
    )

    records: list[Record] = []
    seen_keys: set[str] = set()

    for path in files:
        vendor = _vendor_name_from_file(path)
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            for sheet_name in workbook.sheetnames:
                thread_brand = sheet_name.strip()
                sheet = workbook[sheet_name]
                rows = sheet.iter_rows(values_only=True)
                try:
                    header_row = next(rows)
                except StopIteration:
                    continue

                index_map = _header_index_map(header_row)
                i_name = _col(index_map, "thread name")
                i_code = _col(index_map, "thread number")
                i_pms = _col(index_map, "pms number")
                i_r = _col(index_map, "r")
                i_g = _col(index_map, "g")
                i_b = _col(index_map, "b")

                if i_code == -1:
                    continue  # sheet doesn't match the expected layout

                for row in rows:
                    if i_code >= len(row):
                        continue
                    thread_code = row[i_code]
                    if thread_code is None or str(thread_code).strip() == "":
                        continue

                    thread_name = str(row[i_name]).strip() if i_name != -1 and row[i_name] is not None else ""
                    thread_color_name = thread_name if thread_name and thread_name.upper() != "NA" else ""

                    r = row[i_r] if i_r != -1 and i_r < len(row) else None
                    g = row[i_g] if i_g != -1 and i_g < len(row) else None
                    b = row[i_b] if i_b != -1 and i_b < len(row) else None
                    hex_color = _to_hex(r, g, b)

                    pms_code, pms_name = _normalize_pms(row[i_pms] if i_pms != -1 and i_pms < len(row) else "")

                    key = f"{vendor}|{thread_brand}|{thread_code}|{pms_code}".lower()
                    if key in seen_keys:
                        continue
                    seen_keys.add(key)

                    records.append(
                        Record(
                            vendor=vendor,
                            thread_brand=thread_brand,
                            thread_code=str(thread_code).strip(),
                            thread_color_name=thread_color_name,
                            thread_hex=hex_color,
                            pms_code=pms_code,
                            pms_name=pms_name,
                            # The vendor charts render one RGB swatch per row to
                            # visually represent the assigned PMS match — reused
                            # for both previews since no separately measured
                            # Pantone ink value is provided.
                            pms_hex=hex_color if pms_code else "",
                            updated_at=updated_at,
                        )
                    )
        finally:
            workbook.close()

    return records
